"""Сравнение ответов /summarize с эталоном и классификация причин ошибок.

    python analyze_eval.py                  # все ответы из responses/
    python analyze_eval.py --ids pilot_ids.txt --xlsx natijalar_pilot.xlsx
    python analyze_eval.py --sets 1,2 --xlsx natijalar_jami.xlsx --jsonl results_jami.jsonl

Пишет:
    results.jsonl       — построчно: письмо × поле (эталон, ответ, итог, причина)
    natijalar.xlsx      — Hujjatlar / Xulosa / Xato turlari / Jo‘natuvchilar
"""
import argparse
import collections
import json
import re
import statistics
import unicodedata
from functools import lru_cache
from pathlib import Path

from rapidfuzz import fuzz, process

HERE = Path(__file__).resolve().parent
SETS = {   # набор писем → (папка эталона, папка ответов); id в наборах не пересекаются
    "1": (HERE.parent / "db_test_kirimXat", HERE / "responses"),
    "2": (HERE.parent / "db_test_kirimXat_1000_2", HERE / "responses_2"),
}
REF_DIR, RESP_DIR = SETS["1"]
DIR_DATA = HERE.parent / "doc_summarizer" / "data"   # только чтение


@lru_cache(maxsize=None)
def set_of(doc_id) -> str:
    for name, (ref_dir, _) in SETS.items():
        if (ref_dir / f"DOC_{doc_id}").is_dir():
            return name
    raise FileNotFoundError(f"DOC_{doc_id} нет ни в одном наборе")


def resp_path(doc_id, suffix=".json") -> Path:
    return SETS[set_of(doc_id)][1] / f"{doc_id}{suffix}"

FIELDS = ["number", "date", "signer", "org", "journal", "xdfu", "urgent"]
FIELD_TITLES = {
    "number": "Raqam", "date": "Sana", "signer": "Imzolagan", "org": "Jo‘natuvchi",
    "journal": "Jurnal", "xdfu": "XDFU", "urgent": "Shoshilinch",
}
CONTENT_JOURNALS = {"9", "11", "16"}

# --- итоги (outcome) ------------------------------------------------------------------
OK = "верно"
OK_TRANSLIT = "верно (транслит)"
WRONG = "неверно"
EMPTY_RESP = "пусто у сервиса"
EMPTY_REF = "пусто в эталоне"
FORMAT = "формат"
INITIALS = "инициалы"
INITIALS_PARTIAL = "неполные инициалы"
NOT_PROCESSED = "не обработано"
ORG_VARIANT = "формат / другое написание организации"   # подразделение или имя с вышестоящим
OPERATOR = "решение оператора, в тексте отметки нет"    # срочность: не ошибка сервиса
CORRECT = {OK, OK_TRANSLIT}
WITH_FORMAT = CORRECT | {FORMAT, ORG_VARIANT}          # «с учётом колонки формат»
EXCLUDED = {EMPTY_REF, NOT_PROCESSED, OPERATOR}        # не входят в «всего»
YELLOW = {FORMAT, ORG_VARIANT, INITIALS_PARTIAL, OPERATOR}   # + любые «под сомнением»
ERRORS = {WRONG, EMPTY_RESP, INITIALS}

# --- причины ----------------------------------------------------------------------------
R_OCR = "OCR не прочитал"
R_NOT_EXTRACTED = "прочитал, но не извлёк"
R_BEYOND = "за пределами прочитанного"
R_DOUBT = "эталон под сомнением"
R_NOT_IN_DIR = "нет в справочнике"
R_PARENT = "выбрана вышестоящая"
R_WRONG_HEADER = "не та шапка"
R_GARBAGE = "мусорная запись"
R_NO_ORG = "шапка не найдена"
R_INCOMPLETE = "неполное название"
R_NOT_IN_DOC = "нет в документе"          # нет ни в ocr_text, ни в текстовом слое PDF
R_NO_MARK = "пометки в документе нет"     # ХДФУ/срочно: эталон 1, в тексте метки нет
R_FALSE_MARK = "ложная метка"
R_OTHER_DATE = "взята другая дата из текста"
R_J_CONTENT = "журнал по содержанию (9/11/16)"
R_J_BAD_ORG = "из-за неверного отправителя"
R_J_TABLE = "отправитель верный, журнал по таблице другой"

# --- нормализация -----------------------------------------------------------------------
APOS = "'ʻʼ‘’`´ʹ′‛"
DASHES = "‐‑‒–—―−"
CYR = {
    "а": "a", "б": "b", "в": "v", "г": "g", "ғ": "g'", "д": "d", "е": "e", "ё": "yo",
    "ж": "j", "з": "z", "и": "i", "й": "y", "к": "k", "қ": "q", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ў": "o'",
    "ф": "f", "х": "x", "ҳ": "h", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "sh", "ъ": "'",
    "ь": "", "ы": "i", "э": "e", "ю": "yu", "я": "ya", "і": "i", "ї": "i", "є": "e",
}
CYR_RE = re.compile("[Ѐ-ӿ]")


def translit(s: str) -> str:
    out = []
    prev = " "
    for ch in s:
        lo = ch.lower()
        if lo == "е" and not prev.isalpha():          # «Е» в начале слова → ye
            out.append("ye")
        elif lo in CYR:
            out.append(CYR[lo])
        else:
            out.append(lo)
        prev = ch
    return "".join(out)


def norm(s) -> str:
    """Общая нормализация: регистр, апострофы, тире, кириллица → латиница, пробелы."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s))
    for a in APOS:
        s = s.replace(a, "'")
    for d in DASHES:
        s = s.replace(d, "-")
    s = translit(s.lower())
    return re.sub(r"\s+", " ", s).strip()


def compact(s) -> str:
    return re.sub(r"\s+", "", norm(s))


# --- номер ------------------------------------------------------------------------------
def norm_number(s) -> str:
    s = norm(s)
    s = s.replace("№", " ")
    s = re.sub(r"^\s*(no\.?|n)[\s.-]*(?=\d|$)", " ", s)
    s = re.sub(r"[\s-]*(sonli|son)\s*$", "", s)
    s = re.sub(r"\s+", "", s)
    return s.strip(".,;:")


def ref_empty(v) -> bool:
    return v is None or not str(v).strip() or set(str(v).strip()) <= set("-–—_. ")


# --- ФИО ------------------------------------------------------------------------------------
DIGRAPHS = ("sh", "ch", "o'", "g'", "yo", "yu", "ya", "ye", "ts")


def skel(s: str) -> str:
    """Скелет для вариантов транслитерации: x/h, y/i, ye/e, o'/u, q/k, dj/j, …"""
    s = norm(s)
    s = s.replace("o'", "u").replace("g'", "g").replace("'", "")
    s = s.replace("kh", "x").replace("h", "x")
    for a, b in (("ye", "e"), ("yo", "o"), ("yu", "u"), ("ya", "a"), ("y", "i"),
                 ("dj", "j"), ("dzh", "j"), ("zh", "j"), ("ts", "s"), ("q", "k")):
        s = s.replace(a, b)
    s = re.sub(r"(.)\1+", r"\1", s)
    return s


def consonants(s: str) -> str:
    """Скелет без гласных: Abdullaxanov = Abdullaxonov, Allakuliyv = Allakuliyev."""
    return re.sub(r"[aeiou]", "", skel(s))


def initial_of(word: str) -> str:
    w = word.lower()
    for d in DIGRAPHS:
        if w.startswith(d) and d not in ("ye", "ts"):
            return d
    return w[:1]


def parse_fio(s):
    """→ (фамилия, [инициалы]) или None."""
    s = norm(s)
    if not s:
        return None
    s = re.sub(r"[,;()]", " ", s)
    tokens = [t for t in re.split(r"[.\s]+", s) if t]
    if not tokens:
        return None
    full = [t for t in tokens if len(t) > 2 or (len(t) == 2 and t not in DIGRAPHS and not t.endswith("'"))]
    if not full:
        return None
    # фамилия: слово с типичным окончанием, иначе самое длинное
    sur_re = re.compile(r"(ov|ev|ova|eva|iy|aya|in|ina|yev|yeva|ski|skiy|xon|xo'ja|zoda)$")
    cands = [t for t in full if sur_re.search(t)]
    surname = (cands[0] if len(full) > 1 and cands else max(full, key=len))
    rest = list(tokens)
    rest.remove(surname)
    inits = [initial_of(t) for t in rest]
    return surname, inits


def compare_fio(ref, resp):
    """→ (итог, пометка)"""
    pr, ps = parse_fio(ref), parse_fio(resp)
    if pr is None:
        return EMPTY_REF, ""
    if ps is None:
        return EMPTY_RESP, ""
    (rs, ri), (ss, si) = pr, ps
    ri_s = [skel(x) for x in ri]
    si_s = [skel(x) for x in si]
    vowel_note = ""
    if rs == ss:
        sur = OK
    elif skel(rs) == skel(ss):
        sur = OK_TRANSLIT
    elif (consonants(rs) == consonants(ss) and ri_s and si_s
          and ri_s[: min(len(ri_s), len(si_s))] == si_s[: min(len(ri_s), len(si_s))]):
        # a/o и подобные различия гласных при совпадающих инициалах — вариант транслитерации
        sur, vowel_note = OK_TRANSLIT, "фамилия отличается гласными"
    else:
        return WRONG, ""
    if vowel_note:
        return sur, vowel_note
    if ri_s == si_s or not ri_s:
        return sur, ""
    if si_s and ri_s[: len(si_s)] == si_s and len(si_s) < len(ri_s):
        return INITIALS_PARTIAL, f"эталон {len(ri_s)}, сервис {len(si_s)}"
    if not si_s:
        return INITIALS_PARTIAL, "у сервиса без инициалов"
    if si_s[: len(ri_s)] == ri_s and len(si_s) > len(ri_s):
        return sur, "у сервиса больше инициалов"
    return INITIALS, ""


# --- организация --------------------------------------------------------------------------
ORG_DROP = re.compile(
    r"\b[o0u]\W{0,2}zbekiston\s+respublikasi(ning)?\b|\b(mchj|duk|dm|aj|davlat\s+unitar\s+korxonasi)\b")
# сокращения эталона → полная форма
ORG_ABBR = [
    (r"\bfib\b", "fuqarolik ishlari bo'yicha"), (r"\bjib\b", "jinoyat ishlari bo'yicha"),
    (r"\biiv\b", "ichki ishlar vazirligi"), (r"\biibb\b", "ichki ishlar bosh boshqarmasi"),
    (r"\biib\b", "ichki ishlar boshqarmasi"), (r"\bdxx\b", "davlat xavfsizlik xizmati"),
    (r"\bmib\b", "majburiy ijro byurosi"),
]
# грамматические варианты: viloyati/viloyat, shahri/shahar, tumani/tuman
ORG_GRAM = [(r"\bviloyati\b", "viloyat"), (r"\bshahri\b", "shahar"), (r"\btumani\b", "tuman")]


def norm_org(s) -> str:
    s = norm(s)
    s = re.sub(r"[\"«»“”„]", " ", s)
    s = ORG_DROP.sub(" ", s)
    s = re.sub(r"[.,;:]", " ", s)
    for a, b in ORG_ABBR + ORG_GRAM:
        s = re.sub(a, b, s)
    return re.sub(r"\s+", " ", s).strip()


PERSON_RE = re.compile(r"^\s*([a-z]{1,2}'?\.\s*){1,3}[a-z'-]{3,}\s*$|^[a-z'-]{3,}\s+([a-z]{1,2}'?\.\s*){1,3}$")
PERSON_WORDS = re.compile(r"\b(patent vakili|fuqaro|jismoniy shaxs|advokat)\b")


def ref_is_person(ref) -> bool:
    n = norm(ref)
    return bool(PERSON_RE.match(n) or PERSON_WORDS.search(n))


@lru_cache(maxsize=1)
def directory():
    import pandas as pd
    names = set()
    df = pd.read_parquet(DIR_DATA / "index.parquet", columns=["name"])
    names.update(df["name"].dropna().astype(str))
    m = json.loads((DIR_DATA / "ministries.json").read_text(encoding="utf-8"))
    names.update(m.get("names", []))
    normed = sorted({norm_org(n) for n in names if norm_org(n)})
    return set(normed), normed


def in_directory(ref_org):
    """→ (есть ли, лучшее совпадение, score)"""
    exact, lst = directory()
    n = norm_org(ref_org)
    if n in exact:
        return True, n, 100.0
    best = process.extractOne(n, lst, scorer=fuzz.ratio)
    if best and best[1] >= 95:
        return True, best[0], best[1]
    return False, best[0] if best else "", best[1] if best else 0.0


GARBAGE_RE = re.compile(
    r"(yuzasidan|to'g'risida|munosabati bilan|ma'lum qil|so'ra|iltimos|yuborilmoqda|"
    r"taqdim etiladi|qilamiz|beramiz|\w+(moqda|yapti|gan|di|sin|ngiz)\b.*\w+(moqda|yapti|di|sin)\b)")
DATIVE_RE = re.compile(r"\b\w{4,}(ga|ka|qa)\b")
COPY_RE = re.compile(r"nusxa|kopiya|копия|hurmatli|uvazhaem")


TITLE_RE = re.compile(
    r"\b(a j r i m|ajrim|qaror|karor|hisobot|ariza|ilova|bildirgi|xulosa|ma'lumotnoma|"
    r"dalolatnoma|bayonnoma|opredelenie|reshenie|postanovlenie|prikaz|buyruq|tushuntirish xati|"
    r"normativ|loyiha|nomidan|hal qiluv|xal kiluv)")


def is_title_query(nq: str) -> bool:
    """query начинается с заголовка документа (АЖРИМ, HISOBOT, …), а не с шапки."""
    return bool(TITLE_RE.search(nq[:60]))
PARENT_WORDS = re.compile(
    r"(vazirligi|qo'mitasi|agentligi|prokuraturasi|xizmati|vazirlar mahkamasi|oliy sudi|"
    r"inspeksiyasi|hokimligi|administratsiyasi)$")


def looks_garbage(name: str) -> bool:
    n = norm(name)
    return len(n.split()) > 14 or bool(GARBAGE_RE.search(n)) or n.endswith(".")


# --- даты -----------------------------------------------------------------------------------
UZ_MONTHS = ["yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul", "avgust",
             "sent(?:a|ya)br", "okt(?:a|ya)br", "noyabr", "dekabr"]
RU_MONTHS = ["yanvar", "fevral", "mart", "aprel", "ma", "iyun", "iyul", "avgust",
             "sentyabr", "oktyabr", "noyabr", "dekabr"]


def parse_ref_date(s):
    m = re.match(r"\s*(\d{1,2})\.(\d{1,2})\.(\d{4})", str(s or ""))
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else ""


def date_in_text(iso: str, ctext: str) -> bool:
    """Ищет дату в разных написаниях в compact(ocr_text)."""
    if not iso:
        return False
    y, mo, d = iso.split("-")
    di, mi = int(d), int(mo)
    pats = [rf"(?<!\d)0?{di}[./,-]0?{mi}[./,-]({y}|{y[2:]})(?!\d)"]
    for mon in {UZ_MONTHS[mi - 1], RU_MONTHS[mi - 1]}:
        pats.append(rf"{y}\D{{0,8}}(?<!\d)0?{di}\D{{0,4}}{mon}")
        pats.append(rf"(?<!\d)0?{di}\D{{0,4}}{mon}\w{{0,4}}\D{{0,4}}{y}")
    return any(re.search(p, ctext) for p in pats)


# --- загрузка ------------------------------------------------------------------------------
def load_ref(doc_id):
    d = SETS[set_of(doc_id)][0] / f"DOC_{doc_id}"
    js = next(d.glob("*.json"))
    ref = json.loads(js.read_text(encoding="utf-8"))["rekvizit_document"]
    pdf = next(d.glob("*.pdf"), None)
    return ref, pdf


def load_resp(doc_id):
    meta_p = resp_path(doc_id, ".meta.json")
    meta = json.loads(meta_p.read_text(encoding="utf-8")) if meta_p.exists() else None
    resp_p = resp_path(doc_id)
    resp = json.loads(resp_p.read_text(encoding="utf-8")) if resp_p.exists() else None
    return resp, meta


@lru_cache(maxsize=4)
def pdf_pages_text(pdf: str):
    import fitz
    try:
        with fitz.open(pdf) as doc:
            return [p.get_text() for p in doc]
    except Exception:
        return []


def journal_no(s) -> str:
    m = re.match(r"\s*(\d+)\s*/", str(s or ""))
    return m.group(1) if m else ""


# --- сравнение одного письма --------------------------------------------------------------
def analyze_doc(doc_id):
    ref, pdf = load_ref(doc_id)
    resp, meta = load_resp(doc_id)
    row = {
        "id": doc_id, "set": set_of(doc_id), "pdf": str(pdf) if pdf else "",
        "ref_org": ref.get("hujjat_junatuvchi_tashkilot") or "",
        "ref_journal": journal_no(ref.get("jurnal")),
        "status": (meta or {}).get("status"),
        "server_ms": _int((meta or {}).get("server_ms")),
        "client_ms": (meta or {}).get("client_ms"),
        "attempts": (meta or {}).get("attempts"),
        "n_503": (meta or {}).get("n_503", 0),
        "page_count": None, "pages_used": None, "ocr_used": None,
        "fields": {},
    }
    refvals = {
        "number": ref.get("hujjat_raqami"), "date": ref.get("hujjat_sana"),
        "signer": ref.get("imzolagan_shaxs"), "org": ref.get("hujjat_junatuvchi_tashkilot"),
        "journal": ref.get("jurnal"), "xdfu": ref.get("xdfu"), "urgent": ref.get("shoshilinch"),
    }
    if resp is None:
        for f in FIELDS:
            row["fields"][f] = dict(ref=_s(refvals[f]), resp="", outcome=NOT_PROCESSED,
                                    reason=str(row["status"] or "нет ответа"), note="")
        return row

    row.update(page_count=resp.get("page_count"), pages_used=resp.get("pages_used"),
               ocr_used=resp.get("ocr_used"))
    org = resp.get("organization") or {}
    respvals = {
        "number": resp.get("doc_number"), "date": resp.get("doc_date"),
        "signer": resp.get("signed_fio"), "org": org.get("name") if org.get("matched") else "",
        "journal": resp.get("journal_name"), "xdfu": resp.get("is_xdfu"),
        "urgent": resp.get("is_urgent"),
    }
    ocr = resp.get("ocr_text") or ""
    ctx = Ctx(doc_id, pdf, ocr, resp)

    f = row["fields"]
    f["number"] = cmp_number(refvals["number"], respvals["number"], ctx)
    f["date"] = cmp_date(refvals["date"], respvals["date"], ctx, ref.get("kirish_sana"))
    f["signer"] = cmp_signer(refvals["signer"], respvals["signer"], ctx)
    f["org"] = cmp_org(refvals["org"], respvals["org"], org, ctx)
    f["journal"] = cmp_journal(refvals["journal"], respvals["journal"], f["org"], resp)
    f["xdfu"] = cmp_flag(refvals["xdfu"], respvals["xdfu"], ctx,
                         re.compile(r"xdfu|xdf\b|x\.d\.f\.u|\d+-?xdfu"))
    f["urgent"] = cmp_flag(refvals["urgent"], respvals["urgent"], ctx,
                           re.compile(r"shoshilinch|srochno"), operator_if_no_mark=True)

    # для списков отчёта
    row["org_resp"] = {k: org.get(k) for k in ("name", "id", "score", "source", "query")}
    row["journal_exact"] = resp.get("journal_exact")
    row["journal_resp"] = journal_no(resp.get("journal_name"))
    row["esign"] = bool(ESIGN_RE.search(ctx.n_ocr))
    pr = parse_fio(refvals["signer"])
    if pr:
        sk = skel(pr[0])
        row["ref_surname_in_text"] = sk in ctx.s_ocr or sk in skel(ctx.pdf_text())
    else:
        row["ref_surname_in_text"] = None
    return row


def _int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _s(v):
    return "" if v is None else str(v)


class Ctx:
    def __init__(self, doc_id, pdf, ocr, resp):
        self.id = doc_id
        self.pdf = pdf
        self.ocr = ocr
        self.n_ocr = norm(ocr)
        self.c_ocr = compact(ocr)
        self.s_ocr = skel(ocr)
        self.pages_used = resp.get("pages_used") or 0
        self.page_count = resp.get("page_count") or 0

    def unread_text(self):
        """Текст непрочитанных страниц из текстового слоя PDF (если есть)."""
        if not self.pdf or self.pages_used >= self.page_count:
            return ""
        return "\n".join(pdf_pages_text(str(self.pdf))[self.pages_used:])

    def pdf_text(self):
        return "\n".join(pdf_pages_text(str(self.pdf))) if self.pdf else ""


def generic_reason(ctx: Ctx, find, resp_found=False, ref_close=False):
    """Причина ошибки поля.
    find(text) — есть ли правильное (эталонное) значение в тексте;
    resp_found — значение сервиса ясно стоит в ocr_text;
    ref_close  — в ocr_text есть похожее (искажённое) эталонное значение."""
    if find(ctx.ocr):
        return R_NOT_EXTRACTED
    beyond = ctx.pages_used < ctx.page_count
    if beyond:
        unread = ctx.unread_text()
        if unread and find(unread):
            return R_BEYOND + " (подтверждено текстом PDF)"
    pdf_t = ctx.pdf_text()
    has_layer = len(pdf_t.strip()) > 200
    if has_layer and find(pdf_t):
        return R_OCR + " (есть в текстовом слое PDF)"
    if ref_close:
        return R_OCR + " (прочитано с искажением)"
    if resp_found:
        return R_DOUBT
    if beyond:
        return R_BEYOND
    if has_layer:
        if ESIGN_RE.search(ctx.n_ocr):
            return R_NOT_IN_DOC + " (ЭЦП-документ: реквизит только в ЭЦП/QR)"
        return R_NOT_IN_DOC + " (нет и в текстовом слое PDF)"
    return R_OCR


ESIGN_RE = re.compile(r"elektron raqamli imzo (bilan|orqali) tasdiqlangan|eri bilan tasdiqlangan|"
                      r"ijro\.gov\.uz|qrdocs\.")


def finish(outcome, ref, resp, reason="", note=""):
    return dict(ref=_s(ref), resp=_s(resp), outcome=outcome, reason=reason, note=note)


def cmp_number(ref, resp, ctx):
    if ref_empty(ref):
        return finish(EMPTY_REF, ref, resp)
    if not re.search(r"\d", str(ref)):
        return finish(EMPTY_REF, ref, resp, note="в эталоне не номер")
    rn = norm_number(ref)
    sn = norm_number(resp)

    def find(text):
        return bool(rn) and rn in compact(text).replace("№", "")

    ref_close = len(rn) >= 4 and fuzz.partial_ratio(rn, ctx.c_ocr) >= 80
    if not sn:
        return finish(EMPTY_RESP, ref, resp, generic_reason(ctx, find, False, ref_close))
    if rn == sn:
        return finish(OK, ref, resp)
    if sn.startswith("vmq-") and sn[4:] == rn:
        return finish(FORMAT, ref, resp, note="префикс VMQ-")
    parts = [norm_number(p) for p in re.split(r"[,;()]", str(ref)) if p.strip()]
    if len(parts) > 1 and sn in parts:
        return finish(FORMAT, ref, resp, note="составной номер в эталоне, сервис дал одну часть")
    reason = generic_reason(ctx, find, sn in ctx.c_ocr, ref_close)
    note = ""
    if rn and (rn in sn or sn in rn):
        note = "часть номера"
    elif fuzz.ratio(rn, sn) >= 80:
        note = "похожий номер"
    return finish(WRONG, ref, resp, reason, note)


def cmp_date(ref, resp, ctx, kirish=None):
    out = _cmp_date(ref, resp, ctx)
    reg = parse_ref_date(kirish)
    if out["outcome"] == WRONG and reg and str(resp)[:10] > reg:
        # дата документа не может быть позже регистрации — это ошибка сервиса, не эталона
        if out["reason"] == R_DOUBT:
            out["reason"] = R_OTHER_DATE
        out["note"] = (out["note"] + "; " if out["note"] else "") + "дата сервиса позже регистрации"
    elif out["reason"] == R_DOUBT and reg and parse_ref_date(ref) == reg:
        out["note"] = (out["note"] + "; " if out["note"] else "") + "дата эталона = дата регистрации"
    return out


def _cmp_date(ref, resp, ctx):
    iso = parse_ref_date(ref)
    if not iso:
        return finish(EMPTY_REF, ref, resp)

    def find(text):
        return date_in_text(iso, compact(text))

    if not resp:
        return finish(EMPTY_RESP, ref, resp, generic_reason(ctx, find))
    if str(resp)[:10] == iso:
        return finish(OK, ref, resp)
    reason = generic_reason(ctx, find, date_in_text(str(resp)[:10], ctx.c_ocr))
    note = ""
    try:
        import datetime as dt
        delta = (dt.date.fromisoformat(str(resp)[:10]) - dt.date.fromisoformat(iso)).days
        ry, rm, rd = iso.split("-")
        sy, sm, sd = str(resp)[:10].split("-")
        if (ry, rm, rd) == (sy, sd, sm):
            note = "день/месяц переставлены"
        elif (rm, rd) == (sm, sd):
            note = "другой год"
        else:
            note = f"разница {delta:+d} дн."
    except ValueError:
        pass
    return finish(WRONG, ref, resp, reason, note)


def cmp_signer(ref, resp, ctx):
    outcome, note = compare_fio(ref, resp)
    if outcome in CORRECT or outcome == EMPTY_REF:
        return finish(outcome, ref, resp, note=note)
    pr = parse_fio(ref)
    rs_sk = skel(pr[0]) if pr else ""

    def find(text):
        return bool(rs_sk) and rs_sk in skel(text)

    if outcome in (INITIALS, INITIALS_PARTIAL):
        # фамилия совпала — стоят ли инициалы эталона рядом с фамилией
        inits = "".join(skel(x) for x in pr[1])
        flat = re.sub(r"[\s.]", "", ctx.s_ocr)
        found = bool(inits) and (inits + rs_sk in flat or rs_sk + inits in flat)
        return finish(outcome, ref, resp, R_NOT_EXTRACTED if found else R_OCR, note)
    ps = parse_fio(resp)
    resp_in = bool(ps) and skel(ps[0]) in ctx.s_ocr
    ref_close = len(rs_sk) >= 4 and fuzz.partial_ratio(rs_sk, ctx.s_ocr) >= 85
    if outcome == WRONG and ps:
        r = fuzz.ratio(skel(pr[0]), skel(ps[0]))
        if r >= 85:
            note = f"фамилия очень похожа ({r:.0f})"
    return finish(outcome, ref, resp, generic_reason(ctx, find, resp_in, ref_close), note)


def _present(name_norm, text_norm, thr=88):
    return bool(name_norm) and bool(text_norm) and fuzz.partial_ratio(name_norm, text_norm) >= thr


def _with_parent(rn, sn):
    """Одно название = вышестоящая + другое ("Adliya vazirligi [huzuridagi] Navoiy viloyat …").
    Проверяется в обе стороны: вышестоящая может быть и у сервиса, и у эталона."""
    for short, long_ in ((rn, sn), (sn, rn)):
        if not short or len(long_) <= len(short):
            continue
        tail = long_[-len(short):]
        prefix = long_[: len(long_) - len(short)].strip()
        if fuzz.ratio(short, tail) >= 90 and bool(PARENT_WORDS.search(prefix) or CONNECTOR_RE.search(prefix)):
            return True
    return False


CONNECTOR_RE = re.compile(r"\b(huzuridagi|qoshidagi|tizimidagi|tarkibidagi|ning)$")


def _no_parens(s):
    return re.sub(r"\s+", " ", re.sub(r"\([^)]*\)", " ", s)).strip()


# виды «не та шапка»
HK_COURT = "судебный акт"
HK_REPORT = "отчёт/таблица"
HK_ADDRESSEE = "адресат вместо отправителя"
HK_OTHER = "прочее"
REPORT_RE = re.compile(r"\b(hisobot|loyiha|jadval|ilova|normativ|ma'lumotnoma|reestr|ro'yxat)")
COURT_RE = re.compile(r"\b(a j r i m|ajrim|qaror|karor|opredelenie|reshenie|postanovlenie|nomidan|"
                      r"hal qiluv|xal kiluv|hukm|xukm|prigovor)")


def header_reason(nq):
    head = nq[:80]
    if REPORT_RE.search(head):
        kind = HK_REPORT
    elif COURT_RE.search(head):
        kind = HK_COURT
    elif COPY_RE.search(nq) or DATIVE_RE.search(nq):
        kind = HK_ADDRESSEE
    else:
        kind = HK_OTHER
    return R_WRONG_HEADER + f" ({kind})"


def cmp_org(ref, resp_name, org, ctx):
    rn, sn = norm_org(ref), norm_org(resp_name)
    if not rn:
        return finish(EMPTY_REF, ref, resp_name)
    if ref_is_person(ref):
        return finish(EMPTY_REF, ref, resp_name, note="в эталоне физлицо, а не организация")
    query = org.get("query") or ""
    extra = f"query: {query[:150]}" if query else "query пуст"
    if rn == sn:
        return finish(OK, ref, resp_name)
    if sn and skel(rn) == skel(sn):
        return finish(OK_TRANSLIT, ref, resp_name)
    if sn and _with_parent(rn, sn):
        return finish(ORG_VARIANT, ref, resp_name, note="название в справочнике с вышестоящей")
    rn_np, sn_np = _no_parens(rn), _no_parens(sn)
    if sn and rn_np == sn_np:
        return finish(ORG_VARIANT, ref, resp_name, note="подразделение в скобках")
    if sn and sn_np.startswith(rn_np + " "):
        return finish(ORG_VARIANT, ref, resp_name, note="подразделение организации эталона")
    in_dir, best, best_sc = in_directory(ref)
    head = norm_org(ctx.ocr[:1500])
    nq = norm_org(query)
    ref_in_ocr = _present(rn, head)
    q_matches_ref = _present(rn, nq, 85) or _present(nq, rn, 85)
    title_query = is_title_query(nq)
    dir_note = f"в справочнике ближе всего: {best} ({best_sc:.0f}); " if not in_dir and best_sc >= 80 else ""
    if not sn:
        if not in_dir:
            reason = R_NOT_IN_DIR
        elif not query:
            reason = R_NO_ORG
        elif title_query:
            reason = header_reason(nq)
        else:
            reason = R_NOT_EXTRACTED if ref_in_ocr else R_OCR
        return finish(EMPTY_RESP, ref, resp_name, reason, dir_note + extra)

    ratio = fuzz.ratio(rn, sn)
    note = f"очень похоже ({ratio:.0f}); " if ratio >= 90 else ""
    s_tok, r_tok = set(sn.split()), set(rn.split())
    if looks_garbage(resp_name):
        reason = R_GARBAGE
    elif not in_dir:
        reason = R_NOT_IN_DIR
    elif sn == "adliya vazirligi" and "adliya" not in rn:
        # все письма адресованы Минюсту: вне системы юстиции это адресат, а не вышестоящая
        reason = R_WRONG_HEADER + f" ({HK_ADDRESSEE})"
    elif nq and title_query:
        reason = header_reason(nq)
    elif nq and not q_matches_ref and (COPY_RE.search(nq) or DATIVE_RE.search(nq)):
        reason = R_WRONG_HEADER + f" ({HK_ADDRESSEE})"
    elif nq and not q_matches_ref and _deep_in_text(nq, ctx.n_ocr):
        reason = R_WRONG_HEADER + f" ({HK_OTHER})"
    elif rn.startswith(sn) or (PARENT_WORDS.search(sn) and _present(sn, nq, 85)
                               and (_present(rn, nq, 85) or _present(rn, head, 85))):
        reason = R_PARENT
    elif s_tok < r_tok:
        reason = R_INCOMPLETE
    elif _present(sn, head, 95) and not ref_in_ocr and not _present(rn, norm_org(ctx.pdf_text()[:3000])):
        reason = R_DOUBT
    elif ref_in_ocr:
        reason = R_NOT_EXTRACTED
    else:
        reason = R_OCR
    return finish(WRONG, ref, resp_name, reason, note + dir_note + extra)


def _deep_in_text(nq, n_ocr):
    if not nq or not n_ocr:
        return False
    pos = n_ocr.find(nq[:40])
    return pos > 0 and pos > 0.35 * len(n_ocr)


def cmp_journal(ref, resp, org_field, resp_json):
    rj, sj = journal_no(ref), journal_no(resp)
    if not rj:
        return finish(EMPTY_REF, ref, resp)
    note = f"journal_exact={resp_json.get('journal_exact')}"
    if not sj:
        return finish(EMPTY_RESP, ref, resp, R_J_CONTENT if rj in CONTENT_JOURNALS else "", note)
    if rj == sj:
        return finish(OK, ref, resp, note=note)
    if rj in CONTENT_JOURNALS:
        reason = R_J_CONTENT
    elif org_field["outcome"] not in WITH_FORMAT:
        reason = R_J_BAD_ORG
    else:
        reason = R_J_TABLE
    return finish(WRONG, ref, resp, reason, note)


def cmp_flag(ref, resp, ctx, marker_re, operator_if_no_mark=False):
    if ref is None:
        return finish(EMPTY_REF, ref, resp)
    if resp is None or resp == "":
        return finish(EMPTY_RESP, ref, resp)
    if int(ref) == int(resp):
        return finish(OK, ref, resp)
    m = marker_re.search(ctx.n_ocr)
    ctx_snip = ctx.n_ocr[max(0, m.start() - 40): m.end() + 40] if m else ""
    if int(ref) == 1:
        if m:
            reason, note = R_NOT_EXTRACTED, f"в тексте: …{ctx_snip}…"
        elif operator_if_no_mark and not marker_re.search(norm(ctx.pdf_text())):
            # отметки нет ни в OCR, ни в текстовом слое PDF — решение оператора, не ошибка сервиса
            return finish(OPERATOR, ref, resp)
        elif ctx.pages_used < ctx.page_count:
            reason, note = R_BEYOND, ""
        else:
            # метки нет в тексте: её ставит оператор (сроки, вид документа) или она в штампе
            reason, note = R_NO_MARK, ""
    else:
        reason = R_DOUBT if m else R_FALSE_MARK
        note = f"в тексте: …{ctx_snip}…, эталон 0" if m else ""
    return finish(WRONG, ref, resp, reason, note)


# --- сводки ------------------------------------------------------------------------------
def summarize(rows):
    out = {}
    for f in FIELDS:
        c = collections.Counter(r["fields"][f]["outcome"] for r in rows)
        out[f] = c
    return out


def accuracy(rows, field, pred=lambda r: True):
    """→ (верно строго, верно с учётом «формат», всего).
    Всего = обработанные письма с непустым эталоном, без «решения оператора»."""
    total = ok = with_fmt = 0
    for r in rows:
        if not pred(r):
            continue
        o = r["fields"][field]["outcome"]
        if o in EXCLUDED:
            continue
        total += 1
        ok += o in CORRECT
        with_fmt += o in WITH_FORMAT
    return ok, with_fmt, total


def pct(a, b):
    return f"{100 * a / b:.1f}%" if b else "—"


def is_bad(fv):
    return fv["outcome"] in ERRORS


def is_yellow(fv):
    return fv["outcome"] in YELLOW or fv["reason"] == R_DOUBT or "очень похоже" in fv.get("note", "")


# --- Excel ------------------------------------------------------------------------------
def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    red = PatternFill("solid", fgColor="F8CBAD")
    yellow = PatternFill("solid", fgColor="FFE699")
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    bold = Font(bold=True)
    wb = Workbook()

    # Hujjatlar
    ws = wb.active
    ws.title = "Hujjatlar"
    hdr = ["ID", "To‘plam", "Jo‘natuvchi (etalon)", "Sahifalar", "Server vaqti, s", "Holat"]
    for f in FIELDS:
        t = FIELD_TITLES[f]
        hdr += [f"{t}: etalon", f"{t}: servis", f"{t}: natija", f"{t}: sabab"]
    ws.append(hdr)
    for r in rows:
        pages = f"{r['pages_used']}/{r['page_count']}" if r["page_count"] is not None else ""
        line = [int(r["id"]), int(r["set"]), r["ref_org"], pages,
                round(r["server_ms"] / 1000, 1) if r["server_ms"] else None, str(r["status"])]
        for f in FIELDS:
            v = r["fields"][f]
            reason = v["reason"] + (f" | {v['note']}" if v.get("note") else "")
            line += [v["ref"], v["resp"], v["outcome"], reason]
        ws.append(line)
        i = ws.max_row
        if r["pdf"]:
            c = ws.cell(i, 1)
            c.hyperlink = Path(r["pdf"]).as_uri()
            c.style = "Hyperlink"
        for k, f in enumerate(FIELDS):
            v = r["fields"][f]
            fill = red if is_bad(v) else yellow if is_yellow(v) else None
            if fill:
                for j in range(4):
                    ws.cell(i, 7 + 4 * k + j).fill = fill
    for c in ws[1]:
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    widths = [10, 8, 40, 8, 9, 8] + [22, 22, 14, 30] * len(FIELDS)
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w

    # Xulosa
    ws = wb.create_sheet("Xulosa")
    outcomes = [OK, OK_TRANSLIT, FORMAT, ORG_VARIANT, INITIALS_PARTIAL, INITIALS, WRONG, EMPTY_RESP,
                EMPTY_REF, OPERATOR, NOT_PROCESSED]
    ws.append(["Maydon", "To‘g‘ri (qat’iy)", "To‘g‘ri (format bilan)", "Jami",
               "Aniqlik (qat’iy)", "Aniqlik (format bilan)", "Etalon shubhali"] + outcomes)

    def add(title, field, pred=lambda r: True):
        ok, fmt, tot = accuracy(rows, field, pred)
        c = collections.Counter(r["fields"][field]["outcome"] for r in rows if pred(r))
        doubt = sum(r["fields"][field]["reason"] == R_DOUBT for r in rows if pred(r))
        ws.append([title, ok, fmt, tot, pct(ok, tot), pct(fmt, tot), doubt] + [c.get(o, 0) for o in outcomes])

    xulosa_rows = [(FIELD_TITLES[f], f, lambda r: True) for f in FIELDS] + [
        ("Jurnal (9/11/16 dan tashqari)", "journal", lambda r: r["ref_journal"] not in CONTENT_JOURNALS),
        ("Jurnal (faqat 9/11/16)", "journal", lambda r: r["ref_journal"] in CONTENT_JOURNALS)]
    for title, field, pred in xulosa_rows:
        add(title, field, pred)
    sets = sorted({r["set"] for r in rows})
    if len(sets) > 1:
        # та же точность по каждому набору отдельно и вместе — видно, устойчивы ли цифры
        ws.append([])
        hdr2 = ["To‘plamlar bo‘yicha"]
        for s in sets + [None]:
            name = f"{s}-to‘plam" if s else "Jami"
            hdr2 += [f"{name}: qat’iy", f"{name}: format bilan", f"{name}: soni"]
        ws.append(hdr2)
        for c in ws[ws.max_row]:
            c.font = bold
            c.fill = head_fill
        for title, field, pred in xulosa_rows:
            line = [title]
            for s in sets + [None]:
                ok, fmt, tot = accuracy(rows, field, lambda r, s=s, pred=pred: pred(r) and s in (None, r["set"]))
                line += [pct(ok, tot), pct(fmt, tot), tot]
            ws.append(line)
    ws.append([])
    n_op = sum(r["fields"]["urgent"]["outcome"] == OPERATOR for r in rows)
    ws.append(["Shoshilinch: operator qarori, matnda belgi yo‘q (xato hisoblanmaydi)", n_op])
    es = esign_stats(rows)
    ws.append(["ERI bilan imzolangan hujjatlar", es["esign"]])
    ws.append(["  ulardan imzolovchi familiyasi matnda yo‘q", es["esign_no_surname"]])
    ws.append(["  ulardan servis imzolovchini to‘g‘ri bergan", es["esign_no_surname_ok"]])
    ws.append([])
    st = collections.Counter(str(r["status"]) for r in rows)
    ws.append(["Holatlar (HTTP)"] + [f"{k}: {v}" for k, v in sorted(st.items())])
    ms = sorted(r["server_ms"] for r in rows if r["server_ms"])
    if ms:
        ws.append(["Server vaqti, s", "median", round(statistics.median(ms) / 1000, 1),
                   "p90", round(ms[int(0.9 * (len(ms) - 1))] / 1000, 1)])
    ws.append(["503 javoblar (qayta urinishlar)", sum(r["n_503"] or 0 for r in rows)])
    for c in ws[1]:
        c.font = bold
        c.fill = head_fill
    ws.column_dimensions["A"].width = 32
    for col in "BCDEFGHIJKLMN":
        ws.column_dimensions[col].width = 14

    # Xato turlari
    ws = wb.create_sheet("Xato turlari")
    ws.append(["Maydon", "Natija", "Sabab", "Soni"])
    cnt = collections.Counter()
    for r in rows:
        for f in FIELDS:
            v = r["fields"][f]
            if v["outcome"] not in CORRECT and v["outcome"] not in (EMPTY_REF,):
                cnt[(FIELD_TITLES[f], v["outcome"], reason_key(v["reason"]))] += 1
    for (f, o, rsn), n in sorted(cnt.items(), key=lambda x: (FIELDS.index(_fkey(x[0][0])), -x[1])):
        ws.append([f, o, rsn, n])
    for c in ws[1]:
        c.font = bold
        c.fill = head_fill
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 45
    ws.auto_filter.ref = ws.dimensions

    # Jo‘natuvchilar
    ws = wb.create_sheet("Jo‘natuvchilar")
    groups = collections.defaultdict(list)
    spell = collections.defaultdict(collections.Counter)
    for r in rows:
        k = norm_org(r["ref_org"])
        groups[k].append(r)
        spell[k][r["ref_org"]] += 1
    ws.append(["Jo‘natuvchi", "Xatlar", "Xatolar soni"] + [FIELD_TITLES[f] for f in FIELDS])
    lines = []
    for k, rs in groups.items():
        errs = sum(is_bad(r["fields"][f]) for r in rs for f in FIELDS)
        accs = []
        for f in FIELDS:
            ok, _, tot = accuracy(rs, f)
            accs.append(f"{ok}/{tot}" + (f" ({100 * ok / tot:.0f}%)" if tot else ""))
        lines.append([spell[k].most_common(1)[0][0], len(rs), errs] + accs)
    for l in sorted(lines, key=lambda x: (-x[2], -x[1])):
        ws.append(l)
    for c in ws[1]:
        c.font = bold
        c.fill = head_fill
    ws.column_dimensions["A"].width = 60
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # Jurnal jadvali: отправитель верный, журнал другой
    ws = wb.create_sheet("Jurnal jadvali")
    ws.append(["ID", "Jo‘natuvchi (etalon)", "Jo‘natuvchi (servis, organization.name)", "organization.id",
               "Jo‘natuvchi natija", "journal_exact", "Kutilgan jurnal", "Servis jurnali", "6 o‘rniga 5?"])
    for r in journal_table_cases(rows):
        ws.append([int(r["id"]), r["ref_org"], r["org_resp"]["name"], r["org_resp"]["id"],
                   r["fields"]["org"]["outcome"], r["journal_exact"], r["ref_journal"], r["journal_resp"],
                   "ha" if (r["ref_journal"], r["journal_resp"]) == ("5", "6") else ""])
        ws.cell(ws.max_row, 1).hyperlink = Path(r["pdf"]).as_uri()
    _style_sheet(ws, bold, head_fill, [10, 45, 60, 12, 22, 12, 12, 12, 10])

    # Noto‘g‘ri sarlavha: «не та шапка» по видам
    ws = wb.create_sheet("Noto‘g‘ri sarlavha")
    ws.append(["ID", "Turi", "Jo‘natuvchi (etalon)", "Servis", "query"])
    for kind, r in wrong_header_cases(rows):
        ws.append([int(r["id"]), kind, r["ref_org"], r["fields"]["org"]["resp"],
                   (r["org_resp"]["query"] or "")[:300]])
        ws.cell(ws.max_row, 1).hyperlink = Path(r["pdf"]).as_uri()
    _style_sheet(ws, bold, head_fill, [10, 26, 45, 45, 90])

    # Qisqacha mazmun: ручная проверка 40 писем + автоматические счётчики по всем ответам
    rev_path = HERE / "summary_review.json"
    if rev_path.exists():
        ws = wb.create_sheet("Qisqacha mazmun")
        auto = summary_auto_stats(rows)
        ws.append(["Avtomatik (barcha javoblar)", "Soni", "Jami"])
        ws.append(["Kirill harfi bor summary", auto["cyr"], auto["n"]])
        ws.append(["Asosan kirillda (>50% harf)", auto["mostly"], auto["n"]])
        ws.append(["Jumla oxirida uzilgan (nuqta bilan tugamaydi)", auto["cut"], auto["n"]])
        ws.append([])
        rev = json.loads(rev_path.read_text(encoding="utf-8"))
        ws.append(["ID", "Lotinda", "Xatga zid emas", "Mohiyatni aks ettiradi", "Izoh", "Summary"])
        hdr_row = ws.max_row
        for v in sorted(rev, key=lambda x: x["id"]):
            s = json.loads(resp_path(v["id"]).read_text(encoding="utf-8")).get("summary", "")
            ws.append([int(v["id"]), "ha" if v["lotin"] else "yo‘q", v["zid_emas"], v["mohiyat"], v["izoh"], s])
            if not v["lotin"] or v["zid_emas"] != "ha" or v["mohiyat"] != "ha":
                for j in range(1, 7):
                    ws.cell(ws.max_row, j).fill = yellow
        ws.append([])
        ws.append(["Jami", f"{sum(v['lotin'] for v in rev)}/{len(rev)}",
                   f"{sum(v['zid_emas'] == 'ha' for v in rev)}/{len(rev)}",
                   f"{sum(v['mohiyat'] == 'ha' for v in rev)}/{len(rev)}"])
        for c in list(ws[1]) + list(ws[hdr_row]):
            c.font = bold
            c.fill = head_fill
        for k, w in enumerate([12, 10, 14, 14, 70, 100], 1):
            ws.column_dimensions[get_column_letter(k)].width = w
    wb.save(path)


def summary_auto_stats(rows):
    n = cyr = mostly = cut = 0
    for r in rows:
        p = resp_path(r["id"])
        if not p.exists():
            continue
        s = (json.loads(p.read_text(encoding="utf-8")).get("summary") or "").strip()
        n += 1
        letters = re.findall(r"[^\W\d_]", s)
        c = sum(1 for ch in letters if "Ѐ" <= ch <= "ӿ")
        cyr += c > 0
        mostly += bool(letters) and c / len(letters) > 0.5
        cut += bool(s) and not re.search(r"[.!?…»\"”)]$", s)
    return {"n": n, "cyr": cyr, "mostly": mostly, "cut": cut}


def _style_sheet(ws, bold, head_fill, widths):
    from openpyxl.utils import get_column_letter
    for c in ws[1]:
        c.font = bold
        c.fill = head_fill
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def reason_key(reason):
    """Ключ для группировки причин: «не та шапка» — с видом, остальные — без уточнения в скобках."""
    if not reason:
        return ""
    if reason.startswith(R_WRONG_HEADER):
        return reason
    return reason.split(" (")[0]


def journal_table_cases(rows):
    """Отправитель верный (строго или «формат»), эталонный журнал не 9/11/16, а журнал сервиса другой."""
    out = []
    for r in rows:
        j = r["fields"]["journal"]
        if j["outcome"] == WRONG and j["reason"] == R_J_TABLE:
            out.append(r)
    return sorted(out, key=lambda r: ((r["ref_journal"], r["journal_resp"]) != ("5", "6"), r["ref_org"]))


def wrong_header_cases(rows):
    out = []
    for r in rows:
        rs = r["fields"]["org"]["reason"]
        if rs.startswith(R_WRONG_HEADER):
            m = re.search(r"\((.*)\)", rs)
            out.append((m.group(1) if m else HK_OTHER, r))
    order = [HK_COURT, HK_REPORT, HK_ADDRESSEE, HK_OTHER]
    return sorted(out, key=lambda x: (order.index(x[0]) if x[0] in order else 9, x[1]["id"]))


def esign_stats(rows):
    es = [r for r in rows if r.get("esign")]
    no_sur = [r for r in es if r.get("ref_surname_in_text") is False]
    return {
        "esign": len(es),
        "esign_no_surname": len(no_sur),
        "esign_no_surname_ok": sum(r["fields"]["signer"]["outcome"] in CORRECT for r in no_sur),
        "esign_no_surname_empty": sum(r["fields"]["signer"]["outcome"] == EMPTY_RESP for r in no_sur),
        "no_esign": sum(1 for r in rows if r.get("esign") is False),
        "no_esign_no_surname": sum(1 for r in rows if r.get("esign") is False
                                   and r.get("ref_surname_in_text") is False),
    }


def _fkey(title):
    return next(k for k, v in FIELD_TITLES.items() if v == title)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ids", help="файл со списком id; по умолчанию все письма наборов --sets")
    ap.add_argument("--sets", default="1", help="наборы через запятую: 1, 2 или 1,2")
    ap.add_argument("--xlsx", default="natijalar.xlsx")
    ap.add_argument("--jsonl", default="results.jsonl")
    args = ap.parse_args()
    if args.ids:
        ids = [l.strip() for l in (HERE / args.ids).read_text(encoding="utf-8").split() if l.strip()]
    else:
        ids = [p.name[4:] for s in args.sets.split(",") for p in sorted(SETS[s.strip()][0].glob("DOC_*"))]
    rows = [analyze_doc(i) for i in ids]
    with open(HERE / args.jsonl, "w", encoding="utf-8") as fh:
        for r in rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    write_xlsx(rows, HERE / args.xlsx)
    for f in FIELDS:
        ok, len_, tot = accuracy(rows, f)
        print(f"{FIELD_TITLES[f]:14s} {ok:4d}/{tot:<4d} {pct(ok, tot):>7s}  (с учётом «формат» {pct(len_, tot)})")
    ok, _, tot = accuracy(rows, "journal", lambda r: r["ref_journal"] not in CONTENT_JOURNALS)
    print(f"{'Jurnal без 9/11/16':14s} {ok:4d}/{tot:<4d} {pct(ok, tot):>7s}")
    sets = sorted({r["set"] for r in rows})
    if len(sets) > 1:
        print("\nСтрого по наборам: " + " | ".join(f"набор {s}" for s in sets))
        for f in FIELDS:
            print(f"{FIELD_TITLES[f]:14s} " + " | ".join(
                pct(*accuracy(rows, f, lambda r, s=s: r["set"] == s)[::2]) for s in sets))
    print("->", HERE / args.xlsx)


if __name__ == "__main__":
    main()
