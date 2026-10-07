"""Проверка natijalar.xlsx: листы, размеры, заливки, гиперссылки."""
from openpyxl import load_workbook

wb = load_workbook("natijalar.xlsx")
for ws in wb.worksheets:
    print(f"{ws.title}: {ws.max_row} строк × {ws.max_column} колонок; freeze={ws.freeze_panes}; filter={ws.auto_filter.ref}")
ws = wb["Hujjatlar"]
fills = {}
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for c in row:
        if c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None, "00000000"):
            fills[c.fill.fgColor.rgb] = fills.get(c.fill.fgColor.rgb, 0) + 1
print("заливки:", fills)
print("гиперссылка A2:", ws["A2"].hyperlink.target if ws["A2"].hyperlink else None)
print("шапка:", [c.value for c in ws[1]][:13])
for r in wb["Xulosa"].iter_rows(min_row=1, max_row=16, values_only=True):
    print([x for x in r if x is not None])
